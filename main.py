import customtkinter as ctk
from tkinter import filedialog
import pandas as pd
import smtplib
import re
import os
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from dotenv import load_dotenv
from docx import Document
import openpyxl

load_dotenv()

SMTP_SERVER = os.getenv("SMTP_SERVER")
SMTP_PORT = int(os.getenv("SMTP_PORT"))
SMTP_USER = os.getenv("SMTP_USER")
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD")
USE_TLS = os.getenv("USE_TLS", "True") == "True"

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("dark-blue")


class EmailApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Mail Merge - Excel Sender")
        self.geometry("1000x700")

        self.df = None
        self.file_path = None

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(3, weight=1)

        # Sidebar
        self.sidebar = ctk.CTkFrame(self, width=200)
        self.sidebar.grid(row=0, column=0, rowspan=4, sticky="nsw", padx=10, pady=10)

        self.load_btn = ctk.CTkButton(
            self.sidebar, text="Carica Excel", command=self.load_file
        )
        self.load_btn.pack(pady=10, fill="x")

        self.load_word_btn = ctk.CTkButton(
            self.sidebar, text="Carica Template Word", command=self.load_word_template
        )
        self.load_word_btn.pack(pady=10, fill="x")

        self.create_word_btn = ctk.CTkButton(
            self.sidebar, text="Crea Template Word", command=self.create_word_template
        )
        self.create_word_btn.pack(pady=10, fill="x")

        self.create_excel_btn = ctk.CTkButton(
            self.sidebar, text="Crea Template Excel", command=self.create_excel_template
        )
        self.create_excel_btn.pack(pady=10, fill="x")

        self.preview_btn = ctk.CTkButton(
            self.sidebar, text="Anteprima Prima Email", command=self.preview_email
        )
        self.preview_btn.pack(pady=10, fill="x")

        self.fields_btn = ctk.CTkButton(
            self.sidebar, text="Mostra Campi", command=self.show_fields
        )
        self.fields_btn.pack(pady=10, fill="x")

        self.test_mode = ctk.CTkCheckBox(
            self.sidebar, text="Modalità TEST (invio solo a me)"
        )
        self.test_mode.pack(pady=20)

        self.send_btn = ctk.CTkButton(
            self.sidebar, text="Invia Email", command=self.send_emails
        )
        self.send_btn.pack(pady=20, fill="x")

        # Oggetto
        self.subject_label = ctk.CTkLabel(self, text="Oggetto Email")
        self.subject_label.grid(row=0, column=1, sticky="w", padx=10, pady=(10, 0))

        self.subject_entry = ctk.CTkEntry(self)
        self.subject_entry.grid(row=1, column=1, sticky="ew", padx=10)

        # Corpo
        self.body_label = ctk.CTkLabel(
            self,
            text="Corpo Email (usa {{NomeColonna}} per dichiarare le colonne Excel)",
        )
        self.body_label.grid(row=2, column=1, sticky="nw", padx=10, pady=(10, 0))

        self.body_text = ctk.CTkTextbox(self)
        self.body_text.grid(row=3, column=1, sticky="nsew", padx=10, pady=(0, 10))

        # Progress
        self.progress = ctk.CTkProgressBar(self)
        self.progress.grid(row=4, column=1, sticky="ew", padx=10, pady=5)
        self.progress.set(0)

        self.status_label = ctk.CTkLabel(self, text="Nessun file caricato", anchor="w")
        self.status_label.grid(row=5, column=1, sticky="ew", padx=10, pady=(0, 10))

    # --- Funzioni ---
    def show_dialog(self, title, message, width=500, height=250):
        dialog = ctk.CTkToplevel(self)
        dialog.title(title)
        dialog.geometry(f"{width}x{height}")
        dialog.grab_set()

        dialog.grid_columnconfigure(0, weight=1)
        dialog.grid_rowconfigure(0, weight=1)

        textbox = ctk.CTkTextbox(dialog, wrap="word")
        textbox.grid(row=0, column=0, sticky="nsew", padx=20, pady=20)
        textbox.insert("1.0", message)
        textbox.configure(state="disabled")

        close_btn = ctk.CTkButton(dialog, text="Chiudi", command=dialog.destroy)
        close_btn.grid(row=1, column=0, pady=(0, 15))

    def load_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not path:
            return
        try:
            self.df = pd.read_excel(path)
            self.file_path = path
            self.status_label.configure(text=f"File caricato: {os.path.basename(path)}")
        except Exception as e:
            self.show_dialog("Errore", f"Errore caricamento Excel:\n{e}")

    def load_word_template(self):
        path = filedialog.askopenfilename(filetypes=[("Word files", "*.docx")])
        if not path:
            return
        try:
            doc = Document(path)
            full_text = [
                para.text for para in doc.paragraphs if para.text.strip() != ""
            ]
            if not full_text:
                self.show_dialog("Errore", "Il file Word è vuoto.")
                return
            subject = full_text[0]
            body = "\n\n".join(full_text[1:]) if len(full_text) > 1 else ""

            self.subject_entry.delete(0, "end")
            self.subject_entry.insert(0, subject)

            self.body_text.delete("1.0", "end")
            self.body_text.insert("1.0", body)

            self.show_dialog(
                "Template caricato", "Oggetto e corpo caricati correttamente."
            )
        except Exception as e:
            self.show_dialog("Errore", f"Errore caricamento Word:\n{e}")

    def create_word_template(self):
        try:
            doc = Document()
            doc.add_paragraph("Promemoria pagamento {{AnnoCorso}}")  # Oggetto
            doc.add_paragraph("")  # Riga vuota
            doc.add_paragraph(
                """<p>Gentili Genitori,</p>

            <p>Con la presente desideriamo ricordarvi il pagamento della retta relativa al secondo trimestre del corso di musica {{AnnoCorso}} frequentato da vostro/a figlio/a.</p>

            <p><b>L’importo dovuto è pari a {{Prezzo}} €</b> e può essere versato tramite bonifico bancario:</p>
            <ul>
            <li><b>IBAN: IT25L0863165011066000001528</b></li>
            <li><b>Beneficiario: Associazione Filarmonica Sanvitese</b></li>
            </ul>

            <p>Cogliamo inoltre l’occasione per ricordare che è necessario rinnovare la quota associativa e assicurativa, per un importo complessivo di <u>{{QuotaAssociativa}}</u> €.</p>

            <p>Tali quote dovranno essere versate presso la segreteria nei giorni di mercoledì e venerdì, dalle ore 16.30 alle ore 18.30, con pagamento in contanti entro il mese di {{MesePagamento}} {{AnnoPagamento}}.</p>

            <p>Restiamo a disposizione per eventuali chiarimenti e ringraziamo per la collaborazione.</p>

            <hr>

            <p><b>Cordiali saluti,</b><br>
            <b>Segreteria AFIS</b><br>
            <i>E. Pitton</i></p>

            <p>_____________________________</p>

            <p><b>Filarmonica Sanvitese APS</b><br>
            Piazzale Hermann Zotti, 1<br>
            33078 San Vito Al Tagliamento (PN)<br>
            Tel.: 3396771611<br>
            <a href="https://www.afisanvitese.it">www.afisanvitese.it</a><br>
            <a href="mailto:scuola@afisanvitese.it">scuola@afisanvitese.it</a></p>"""
            )

            save_path = filedialog.asksaveasfilename(
                defaultextension=".docx",
                filetypes=[("Word files", "*.docx")],
                initialfile="Matrice_mail.docx",
            )
            if save_path:
                doc.save(save_path)
                self.show_dialog("Creato", f"Template Word salvato come:\n{save_path}")
        except Exception as e:
            self.show_dialog("Errore", f"Errore creazione Word:\n{e}")

    def create_excel_template(self):
        try:
            # Crea workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Template"

            # Intestazioni
            ws.append(["Email", "Nome", "Cognome", "AltroCampo"])

            # Riga esempio
            ws.append(["esempio@email.com", "Mario", "Rossi", "Valore"])

            save_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile="Matrice_mail.xlsx",
            )
            if save_path:
                wb.save(save_path)
                self.show_dialog("Creato", f"Template Excel salvato come:\n{save_path}")
        except Exception as e:
            self.show_dialog("Errore", f"Errore creazione Excel:\n{e}")

    def show_fields(self):
        if self.df is None:
            self.show_dialog("Attenzione", "Caricare prima un file Excel.")
            return
        columns = self.df.columns.tolist()[1:]
        fields = "\n".join([f"{{{{{col}}}}}" for col in columns])
        self.show_dialog("Campi disponibili", fields)

    def validate_placeholders(self, text):
        placeholders = re.findall(r"\{\{(.*?)\}\}", text)
        columns = self.df.columns.tolist()
        for ph in placeholders:
            if ph not in columns:
                return False, ph
        return True, None

    def preview_email(self):
        if self.df is None:
            self.show_dialog("Errore", "Caricare prima un file Excel.")
            return
        subject = self.subject_entry.get()
        body = self.body_text.get("1.0", "end")
        row = self.df.iloc[0]
        for col in self.df.columns:
            subject = subject.replace(f"{{{{{col}}}}}", str(row[col]))
            body = body.replace(f"{{{{{col}}}}}", str(row[col]))
        self.show_dialog(
            "Anteprima Prima Email",
            f"OGGETTO:\n{subject}\n\nCORPO:\n{body}",
            width=600,
            height=400,
        )

    def send_emails(self):
        if self.df is None:
            self.show_dialog("Errore", "Caricare prima un file Excel.")
            return

        subject_template = self.subject_entry.get().strip()
        body_template = self.body_text.get("1.0", "end").strip()
        if not subject_template or not body_template:
            self.show_dialog("Errore", "Oggetto e corpo obbligatori.")
            return

        valid, wrong_field = self.validate_placeholders(
            subject_template + body_template
        )
        if not valid:
            self.show_dialog("Errore Placeholder", f"Campo non trovato: {wrong_field}")
            return

        try:
            server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
            server.ehlo()
            if USE_TLS:
                server.starttls()
                server.ehlo()
            server.login(SMTP_USER, SMTP_PASSWORD)
        except Exception as e:
            self.show_dialog("Errore SMTP", f"Errore di connessione:\n{e}")
            return

        self.send_btn.configure(state="disabled")
        total = len(self.df)
        sent = 0
        errors = 0

        with open("email_log.txt", "a", encoding="utf-8") as log:
            for index, row in self.df.iterrows():
                recipient = (
                    SMTP_USER if self.test_mode.get() else str(row.iloc[0]).strip()
                )
                if not re.match(r"[^@]+@[^@]+\.[^@]+", recipient):
                    errors += 1
                    log.write(f"Email non valida: {recipient}\n")
                    continue

                subject = subject_template
                body = body_template
                for col in self.df.columns:
                    subject = subject.replace(f"{{{{{col}}}}}", str(row[col]))
                    body = body.replace(f"{{{{{col}}}}}", str(row[col]))

                msg = MIMEMultipart()
                msg["From"] = SMTP_USER
                msg["To"] = recipient
                msg["Subject"] = subject
                msg.attach(MIMEText(body, "html"))

                try:
                    server.send_message(msg)
                    sent += 1
                except Exception as e:
                    errors += 1
                    log.write(f"Errore invio a {recipient}: {str(e)}\n")

                self.progress.set((index + 1) / total)
                self.update_idletasks()

        server.quit()
        self.send_btn.configure(state="normal")
        self.show_dialog(
            "Invio completato",
            f"Email inviate: {sent}\nErrori: {errors}",
            width=550,
            height=300,
        )


if __name__ == "__main__":
    app = EmailApp()
    app.mainloop()
